"""
pdf -> text
docx -> text
txt -> text
"""

from pathlib import Path

from docx import Document
from fastapi import HTTPException
from PyPDF2 import PdfReader


def extract_pdf_text(file_path: str | Path) -> str:
    reader = PdfReader(str(file_path))

    text = ""

    for page in reader.pages:
        text += page.extract_text()

    return text


def extract_txt_text(file_path: str | Path) -> str:
    with open(file_path, "r") as file:
        content = file.read()

    return content


def extract_docx_text(file_path: str | Path) -> str:
    document = Document(str(file_path))

    parts = []

    # Paragraph text, in document order
    for paragraph in document.paragraphs:
        if paragraph.text.strip():
            parts.append(paragraph.text)

    # Table cell text - contracts frequently put schedules/fee tables here,
    # and paragraph iteration alone skips them entirely.
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                if cell.text.strip():
                    parts.append(cell.text)

    return "\n".join(parts)


def extract_text(file_path: str | Path) -> str:
    file_path = Path(file_path)

    ext = file_path.suffix.lower()

    if ext == ".pdf":
        return extract_pdf_text(file_path)
    elif ext == ".txt":
        return extract_txt_text(file_path)
    elif ext == ".docx":
        return extract_docx_text(file_path)
    else:
        raise HTTPException(
            status_code=400,
            detail="Unsupported file type. Only .pdf, .docx and .txt allowed!",
        )
